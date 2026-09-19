"""Safe text extraction for the supported enterprise document formats."""
from __future__ import annotations

from dataclasses import dataclass
import mimetypes
from pathlib import Path
import re

from docx import Document as DocxDocument
from pypdf import PdfReader


SUPPORTED_EXTENSIONS = {".md", ".markdown", ".txt", ".pdf", ".docx"}


@dataclass(frozen=True)
class ParsedSection:
    heading: str
    text: str
    page_number: int | None = None


@dataclass(frozen=True)
class ParsedDocument:
    title: str
    mime_type: str
    sections: tuple[ParsedSection, ...]


def parse_document(path: Path, *, title: str = "", max_bytes: int,
                   max_characters: int = 2_000_000, max_pages: int = 500) -> ParsedDocument:
    path = path.resolve()
    if not path.is_file():
        raise ValueError("文档不存在或不是文件")
    if path.suffix.lower() not in SUPPORTED_EXTENSIONS:
        raise ValueError("仅支持 Markdown、TXT、PDF 和 DOCX")
    if path.stat().st_size > max_bytes:
        raise ValueError("文档超过允许的大小")
    mime_type = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    suffix = path.suffix.lower()
    if suffix in {".md", ".markdown"}:
        sections = _parse_markdown(path)
    elif suffix == ".txt":
        sections = (ParsedSection("", path.read_text(encoding="utf-8-sig")),)
    elif suffix == ".pdf":
        sections = _parse_pdf(path, max_pages=max_pages)
    else:
        sections = _parse_docx(path)
    sections = tuple(section for section in sections if section.text.strip())
    if not sections:
        raise ValueError("文档没有可索引的文本内容")
    if sum(len(section.text) for section in sections) > max_characters:
        raise ValueError("文档提取文本超过允许的长度")
    return ParsedDocument(title.strip() or path.stem, mime_type, sections)


def _parse_markdown(path: Path) -> tuple[ParsedSection, ...]:
    sections: list[ParsedSection] = []
    heading = ""
    body: list[str] = []
    for raw in path.read_text(encoding="utf-8-sig").splitlines():
        match = re.match(r"^#{1,6}\s+(.+?)\s*$", raw)
        if match:
            if body:
                sections.append(ParsedSection(heading, "\n".join(body).strip()))
            heading, body = match.group(1).strip(), []
        else:
            body.append(raw)
    if body:
        sections.append(ParsedSection(heading, "\n".join(body).strip()))
    return tuple(sections)


def _parse_pdf(path: Path, *, max_pages: int) -> tuple[ParsedSection, ...]:
    reader = PdfReader(str(path), strict=False)
    if reader.is_encrypted:
        try:
            unlocked = reader.decrypt("")
        except Exception as exc:  # noqa: BLE001 - parser boundary
            raise ValueError("不支持加密 PDF") from exc
        if not unlocked:
            raise ValueError("不支持加密 PDF")
    if len(reader.pages) > max_pages:
        raise ValueError("PDF 页数超过允许的上限")
    return tuple(
        ParsedSection(f"第 {number} 页", page.extract_text() or "", number)
        for number, page in enumerate(reader.pages, 1)
    )


def _parse_docx(path: Path) -> tuple[ParsedSection, ...]:
    document = DocxDocument(str(path))
    sections: list[ParsedSection] = []
    heading = ""
    body: list[str] = []
    for paragraph in document.paragraphs:
        text = paragraph.text.strip()
        if not text:
            continue
        if paragraph.style and paragraph.style.name.lower().startswith("heading"):
            if body:
                sections.append(ParsedSection(heading, "\n".join(body)))
            heading, body = text, []
        else:
            body.append(text)
    if body:
        sections.append(ParsedSection(heading, "\n".join(body)))
    return tuple(sections)
