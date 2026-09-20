"""Structured office artifact generation for the independent admin service."""
from __future__ import annotations

from datetime import datetime, timezone
import json
import os
from pathlib import Path
import re
import tempfile
from typing import Any
from xml.sax.saxutils import escape


SUPPORTED_ARTIFACT_FORMATS = frozenset({"docx", "pdf", "pptx", "xlsx"})
ARTIFACT_MEDIA_TYPES = {
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "pdf": "application/pdf",
    "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
}
MAX_INPUT_CHARACTERS = 300_000
MAX_COLLECTION_ITEMS = 100
MAX_TABLE_ROWS = 2_000


class ArtifactError(ValueError):
    """A validation or rendering failure safe to return to an administrator."""


class ArtifactService:
    """Generate deterministic office files without arbitrary code execution."""

    def __init__(self, output_path: str | Path):
        self.output_path = Path(output_path).resolve()

    def initialize(self) -> None:
        self.output_path.mkdir(parents=True, exist_ok=True)

    def healthcheck(self) -> tuple[bool, str]:
        try:
            self.initialize()
        except OSError:
            return False, "artifact_storage_unavailable"
        if not os.access(self.output_path, os.R_OK | os.W_OK):
            return False, "artifact_storage_not_writable"
        return True, "ok"

    def list(self) -> list[dict]:
        try:
            files = [path for path in self.output_path.iterdir()
                     if path.is_file() and path.suffix.lower().lstrip(".") in SUPPORTED_ARTIFACT_FORMATS]
        except FileNotFoundError:
            return []
        files.sort(key=lambda path: (path.stat().st_mtime_ns, path.name), reverse=True)
        return [self._metadata(path) for path in files]

    def resolve(self, filename: str) -> Path:
        safe = self._filename(filename, Path(filename).suffix.lower().lstrip("."))
        target = (self.output_path / safe).resolve()
        if target.parent != self.output_path or not target.is_file():
            raise ArtifactError("产物不存在")
        return target

    def create(self, payload: dict[str, Any]) -> dict:
        serialized = json.dumps(payload, ensure_ascii=False)
        if len(serialized) > MAX_INPUT_CHARACTERS:
            raise ArtifactError("产物内容超过允许的大小")
        artifact_format = self._text(payload.get("format"), 10).strip().lower().lstrip(".")
        if artifact_format not in SUPPORTED_ARTIFACT_FORMATS:
            raise ArtifactError("format 必须是 docx、pdf、pptx 或 xlsx")

        self.initialize()
        target = self._unique_path(self._filename(payload.get("filename"), artifact_format))
        handle = tempfile.NamedTemporaryFile(
            prefix=".docmind-", suffix=f".{artifact_format}", dir=self.output_path, delete=False,
        )
        temporary = Path(handle.name)
        handle.close()
        try:
            validation = _RENDERERS[artifact_format](payload, temporary)
            os.replace(temporary, target)
        except ArtifactError:
            raise
        except (OSError, ImportError, ValueError) as exc:
            raise ArtifactError(str(exc)) from exc
        except Exception as exc:  # Keep internal details out of HTTP responses.
            raise ArtifactError(f"{artifact_format.upper()} 生成失败") from exc
        finally:
            try:
                temporary.unlink(missing_ok=True)
            except OSError:
                pass
        return {**self._metadata(target), "validation": validation}

    def _metadata(self, path: Path) -> dict:
        stat = path.stat()
        artifact_format = path.suffix.lower().lstrip(".")
        return {
            "filename": path.name,
            "format": artifact_format,
            "bytes": stat.st_size,
            "created_at": datetime.fromtimestamp(stat.st_mtime, timezone.utc).isoformat(),
            "download_url": f"/api/admin/artifacts/{path.name}",
        }

    def _filename(self, value: Any, artifact_format: str) -> str:
        if artifact_format not in SUPPORTED_ARTIFACT_FORMATS:
            raise ArtifactError("产物格式无效")
        raw = self._text(value, 128).strip() or f"document.{artifact_format}"
        if Path(raw).name != raw:
            raise ArtifactError("filename 只能是文件名，不能包含目录")
        stem = Path(raw).stem if Path(raw).suffix else raw
        stem = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", stem).strip(" .")
        if not stem:
            raise ArtifactError("filename 无效")
        return f"{stem}.{artifact_format}"

    def _unique_path(self, filename: str) -> Path:
        target = self.output_path / filename
        if not target.exists():
            return target
        for index in range(2, 10_000):
            candidate = self.output_path / f"{target.stem}-{index}{target.suffix}"
            if not candidate.exists():
                return candidate
        raise ArtifactError("无法分配输出文件名")

    @staticmethod
    def _text(value: Any, limit: int = 20_000) -> str:
        return "" if value is None else str(value)[:limit]


def _text(value: Any, limit: int = 20_000) -> str:
    return "" if value is None else str(value)[:limit]


def _items(value: Any, name: str, limit: int = MAX_COLLECTION_ITEMS) -> list:
    if value is None:
        return []
    if not isinstance(value, list):
        raise ArtifactError(f"{name} 必须是数组")
    if len(value) > limit:
        raise ArtifactError(f"{name} 最多允许 {limit} 项")
    return value


def _sections(payload: dict) -> list[dict]:
    sections = _items(payload.get("sections"), "sections")
    if not all(isinstance(item, dict) for item in sections):
        raise ArtifactError("sections 的每一项必须是对象")
    return sections


def _render_docx(payload: dict, path: Path) -> dict:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml.ns import qn
    from docx.shared import Cm, Pt

    document = Document()
    document.sections[0].top_margin = Cm(2.2)
    document.sections[0].bottom_margin = Cm(2.2)
    document.sections[0].left_margin = Cm(2.4)
    document.sections[0].right_margin = Cm(2.4)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")

    title = _text(payload.get("title"), 512).strip()
    if title:
        heading = document.add_heading(title, level=0)
        heading.alignment = WD_ALIGN_PARAGRAPH.CENTER
    subtitle = _text(payload.get("subtitle"), 1_024).strip()
    if subtitle:
        paragraph = document.add_paragraph(subtitle)
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    def add_block(block: dict) -> None:
        heading = _text(block.get("heading"), 512).strip()
        if heading:
            level = max(1, min(3, int(block.get("level", 1))))
            document.add_heading(heading, level=level)
        for value in _items(block.get("paragraphs"), "paragraphs"):
            document.add_paragraph(_text(value))
        for value in _items(block.get("bullets"), "bullets"):
            document.add_paragraph(_text(value), style="List Bullet")
        if block.get("table"):
            _docx_table(document, block["table"])

    add_block(payload)
    for section in _sections(payload):
        add_block(section)
    document.save(path)
    reopened = Document(path)
    return {"paragraphs": len(reopened.paragraphs), "tables": len(reopened.tables)}


def _docx_table(document, spec: Any) -> None:
    headers, rows = _table_data(spec)
    width = len(headers) or max((len(row) for row in rows), default=0)
    if not width:
        return
    table = document.add_table(rows=1 if headers else 0, cols=width)
    table.style = "Table Grid"
    if headers:
        for index, value in enumerate(headers[:width]):
            table.rows[0].cells[index].text = _text(value)
    for values in rows:
        cells = table.add_row().cells
        for index, value in enumerate(values[:width]):
            cells[index].text = _text(value)


def _register_pdf_font() -> str:
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont

    for font_path in (
        r"C:\Windows\Fonts\msyh.ttc",
        r"C:\Windows\Fonts\simhei.ttf",
        "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
    ):
        if os.path.isfile(font_path):
            try:
                pdfmetrics.registerFont(TTFont("DocMindCJK", font_path))
                return "DocMindCJK"
            except Exception:
                continue
    return "Helvetica"


def _render_pdf(payload: dict, path: Path) -> dict:
    from pypdf import PdfReader
    from reportlab.lib import colors
    from reportlab.lib.enums import TA_CENTER
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import ListFlowable, ListItem, Paragraph, SimpleDocTemplate, Spacer

    font = _register_pdf_font()
    base = getSampleStyleSheet()["BodyText"]
    body = ParagraphStyle("DocMindBody", parent=base, fontName=font, fontSize=10.5,
                          leading=16, wordWrap="CJK", spaceAfter=5)
    title_style = ParagraphStyle("DocMindTitle", parent=body, fontSize=22, leading=28,
                                 alignment=TA_CENTER, spaceAfter=10)
    subtitle_style = ParagraphStyle("DocMindSubtitle", parent=body, fontSize=11,
                                    textColor=colors.HexColor("#555555"),
                                    alignment=TA_CENTER, spaceAfter=14)
    headings = {
        1: ParagraphStyle("DocMindH1", parent=body, fontSize=16, leading=22,
                          spaceBefore=12, spaceAfter=6),
        2: ParagraphStyle("DocMindH2", parent=body, fontSize=13, leading=19,
                          spaceBefore=9, spaceAfter=4),
        3: ParagraphStyle("DocMindH3", parent=body, fontSize=11, leading=17,
                          spaceBefore=7, spaceAfter=3),
    }
    story = []
    title = _text(payload.get("title"), 512).strip()
    if title:
        story.append(Paragraph(escape(title), title_style))
    subtitle = _text(payload.get("subtitle"), 1_024).strip()
    if subtitle:
        story.append(Paragraph(escape(subtitle), subtitle_style))

    def add_block(block: dict) -> None:
        heading = _text(block.get("heading"), 512).strip()
        if heading:
            level = max(1, min(3, int(block.get("level", 1))))
            story.append(Paragraph(escape(heading), headings[level]))
        for value in _items(block.get("paragraphs"), "paragraphs"):
            story.append(Paragraph(escape(_text(value)).replace("\n", "<br/>"), body))
        bullets = _items(block.get("bullets"), "bullets")
        if bullets:
            story.append(ListFlowable(
                [ListItem(Paragraph(escape(_text(value)), body)) for value in bullets],
                bulletType="bullet", leftIndent=16,
            ))
            story.append(Spacer(1, 4))
        if block.get("table"):
            story.append(_pdf_table(block["table"], body))
            story.append(Spacer(1, 7))

    add_block(payload)
    for section in _sections(payload):
        add_block(section)
    document = SimpleDocTemplate(
        str(path), pagesize=A4, rightMargin=18 * mm, leftMargin=18 * mm,
        topMargin=18 * mm, bottomMargin=18 * mm, title=title,
    )
    document.build(story or [Paragraph(" ", body)])
    return {"pages": len(PdfReader(str(path)).pages)}


def _pdf_table(spec: Any, body_style):
    from reportlab.lib import colors
    from reportlab.platypus import Paragraph, Table, TableStyle

    headers, rows = _table_data(spec)
    matrix = []
    if headers:
        matrix.append([Paragraph(escape(_text(value)), body_style) for value in headers])
    matrix.extend(
        [Paragraph(escape(_text(value)), body_style) for value in row]
        for row in rows
    )
    table = Table(matrix or [[Paragraph(" ", body_style)]],
                  repeatRows=1 if headers else 0, hAlign="LEFT")
    table.setStyle(TableStyle([
        ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#B8BEC6")),
        ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#EEF1F4")),
        ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ("LEFTPADDING", (0, 0), (-1, -1), 6),
        ("RIGHTPADDING", (0, 0), (-1, -1), 6),
    ]))
    return table


def _render_pptx(payload: dict, path: Path) -> dict:
    from pptx import Presentation
    from pptx.util import Inches, Pt

    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)
    title = _text(payload.get("title"), 512).strip()
    if title:
        slide = presentation.slides.add_slide(presentation.slide_layouts[0])
        slide.shapes.title.text = title
        slide.placeholders[1].text = _text(payload.get("subtitle"), 1_024)
    slides = _items(payload.get("slides"), "slides")
    if not slides:
        slides = [{
            "title": section.get("heading", ""),
            "bullets": section.get("bullets") or section.get("paragraphs") or [],
        } for section in _sections(payload)]
    for spec in slides:
        if not isinstance(spec, dict):
            raise ArtifactError("slides 的每一项必须是对象")
        slide = presentation.slides.add_slide(presentation.slide_layouts[1])
        slide.shapes.title.text = _text(spec.get("title"), 512)
        frame = slide.placeholders[1].text_frame
        frame.clear()
        bullets = _items(spec.get("bullets"), "slides.bullets")
        if not bullets and spec.get("body"):
            bullets = [spec["body"]]
        for index, value in enumerate(bullets):
            paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
            paragraph.text = _text(value)
            paragraph.level = 0
            paragraph.font.size = Pt(24)
            paragraph.font.name = "Microsoft YaHei"
    if not presentation.slides:
        presentation.slides.add_slide(presentation.slide_layouts[6])
    presentation.save(path)
    return {"slides": len(Presentation(path).slides)}


def _render_xlsx(payload: dict, path: Path) -> dict:
    from openpyxl import Workbook, load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    workbook = Workbook()
    workbook.remove(workbook.active)
    sheets = _items(payload.get("sheets"), "sheets")
    if not sheets:
        table = payload.get("table") or {}
        sheets = [{"name": payload.get("title") or "Data",
                   "headers": table.get("headers", []), "rows": table.get("rows", [])}]
    existing_names: set[str] = set()
    for spec in sheets:
        if not isinstance(spec, dict):
            raise ArtifactError("sheets 的每一项必须是对象")
        name = re.sub(r"[\\/*?:\[\]]", "_", _text(spec.get("name"), 31).strip()) or "Sheet"
        if name in existing_names:
            raise ArtifactError("工作表名称不能重复")
        existing_names.add(name)
        sheet = workbook.create_sheet(name)
        headers = _items(spec.get("headers"), "sheets.headers")
        rows = _items(spec.get("rows"), "sheets.rows", MAX_TABLE_ROWS)
        if headers:
            sheet.append([_xlsx_value(value, payload) for value in headers])
            for cell in sheet[1]:
                cell.font = Font(bold=True, color="FFFFFF")
                cell.fill = PatternFill("solid", fgColor="334155")
                cell.alignment = Alignment(vertical="center")
            sheet.freeze_panes = "A2"
        for row in rows:
            if not isinstance(row, list):
                raise ArtifactError("sheets.rows 的每一项必须是数组")
            sheet.append([_xlsx_value(value, payload) for value in row])
        for column in sheet.columns:
            width = min(60, max(10, max(
                (len(_text(cell.value, 200)) for cell in column), default=0,
            ) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width
    workbook.save(path)
    reopened = load_workbook(path, read_only=True, data_only=False)
    result = {"sheets": reopened.sheetnames}
    reopened.close()
    return result


def _xlsx_value(value: Any, payload: dict) -> Any:
    if isinstance(value, (int, float, bool)) or value is None:
        return value
    text = _text(value)
    if not payload.get("allow_formulas") and text.startswith(("=", "+", "-", "@")):
        return "'" + text
    return text


def _table_data(spec: Any) -> tuple[list, list[list]]:
    if not isinstance(spec, dict):
        raise ArtifactError("table 必须是对象")
    headers = _items(spec.get("headers"), "table.headers")
    rows = _items(spec.get("rows"), "table.rows", MAX_TABLE_ROWS)
    if not all(isinstance(row, list) for row in rows):
        raise ArtifactError("table.rows 的每一项必须是数组")
    return headers, rows


_RENDERERS = {
    "docx": _render_docx,
    "pdf": _render_pdf,
    "pptx": _render_pptx,
    "xlsx": _render_xlsx,
}
