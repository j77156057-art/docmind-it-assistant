from pathlib import Path
import tempfile
import unittest

import httpx
from docx import Document as DocxDocument
from pypdf import PdfWriter

from assistant import ITQueryService
from backend import HybridRetriever, ModelRouter, Principal, QueryDatabase
from backend.embeddings import EmbeddingClient, EmbeddingError
from ingestion import DocumentIngestionService
from ingestion.parsers import parse_document


class DocumentIngestionRetrievalTests(unittest.TestCase):
    def make_database(self, root: str) -> QueryDatabase:
        database = QueryDatabase(str(Path(root) / "knowledge.db"))
        database.initialize()
        return database

    def make_ingestion(self, database: QueryDatabase, embeddings: EmbeddingClient):
        return DocumentIngestionService(
            database, embeddings,
            max_bytes=2 * 1024 * 1024,
            chunk_max_chars=300,
            chunk_overlap_chars=40,
        )

    def test_postgres_lexical_query_uses_safe_or_terms(self):
        value = QueryDatabase._postgres_websearch_query("VPN 凭据 wi-fi")
        self.assertEqual(value, '"vpn" OR "wi-fi" OR "凭据"')

    def test_markdown_import_duplicate_version_and_hybrid_retrieval(self):
        with tempfile.TemporaryDirectory() as root:
            project = Path(root)
            path = project / "vpn.md"
            path.write_text(
                "# 企业网络\n## VPN 凭据过期\n请在统一身份门户更新密码，然后重新登录 VPN。\n",
                encoding="utf-8",
            )
            database = self.make_database(root)
            embeddings = EmbeddingClient(
                mode="hash", provider="builtin", model="hash-1024",
                base_url="", api_key="",
            )
            ingestion = self.make_ingestion(database, embeddings)
            first = ingestion.import_file(path, source_key="handbook/vpn")
            duplicate = ingestion.import_file(path, source_key="handbook/vpn")

            retriever = HybridRetriever(database, embeddings, top_k=3)
            service = ITQueryService(
                str(project / "missing.md"), database, ModelRouter("knowledge"),
                retriever=retriever,
            )
            result = service.query("session-doc", "VPN 凭据过期怎么处理")

            path.write_text(
                "# 企业网络\n## VPN 凭据过期\n请先更新统一身份密码，再等待五分钟登录 VPN。\n",
                encoding="utf-8",
            )
            second = ingestion.import_file(path, source_key="handbook/vpn")
            versions = database.list_documents()
            database.dispose()

        self.assertEqual(first["status"], "indexed")
        self.assertTrue(duplicate["duplicate"])
        self.assertEqual(second["version"], 2)
        self.assertEqual(result["evidence"], "sufficient")
        self.assertIn("统一身份门户", result["answer"])
        self.assertEqual(result["citations"][0]["source"], "vpn")
        self.assertEqual(result["citations"][0]["version"], 1)
        self.assertEqual({item["status"] for item in versions}, {"indexed", "superseded"})

    def test_docx_parser_preserves_heading(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "manual.docx"
            document = DocxDocument()
            document.add_heading("打印机故障", level=1)
            document.add_paragraph("清除卡纸后重新启动打印机。")
            document.save(path)
            parsed = parse_document(path, max_bytes=1024 * 1024)
        self.assertEqual(parsed.sections[0].heading, "打印机故障")
        self.assertIn("清除卡纸", parsed.sections[0].text)

    def test_knowledge_overview_respects_document_acl(self):
        with tempfile.TemporaryDirectory() as root:
            project = Path(root)
            public_path = project / "public.md"
            public_path.write_text("# 公共手册\n## VPN 使用\n连接说明。\n", encoding="utf-8")
            private_path = project / "private.md"
            private_path.write_text("# 私密手册\n## 薪酬系统\n内部说明。\n", encoding="utf-8")
            database = self.make_database(root)
            embeddings = EmbeddingClient(
                mode="hash", provider="builtin", model="hash-1024",
                base_url="", api_key="",
            )
            ingestion = self.make_ingestion(database, embeddings)
            ingestion.import_file(
                public_path, title="公共手册", source_key="public", access_scope="public",
            )
            private = ingestion.import_file(
                private_path, title="私密手册", source_key="private", access_scope="restricted",
            )
            service = ITQueryService(
                str(project / "missing.md"), database,
                ModelRouter("local", local_provider="ollama", local_model="qwen2.5:7b"),
                retriever=HybridRetriever(database, embeddings),
            )
            principal = Principal("viewer-1", frozenset({"viewer"}), frozenset())

            result = service.query("overview", "文档库有哪些内容？", principal)
            database.set_document_acl(
                private["document_id"], [("user", "viewer-1")],
                actor_subject_id="admin-1", request_id="overview-acl",
            )
            authorized = service.query("overview-authorized", "知识库包含哪些内容？", principal)
            database.dispose()

        self.assertEqual(result["model"]["route"], "knowledge")
        self.assertIn("公共手册", result["answer"])
        self.assertIn("VPN 使用", result["answer"])
        self.assertNotIn("私密手册", result["answer"])
        self.assertEqual([item["source"] for item in result["citations"]], ["公共手册"])
        self.assertIn("私密手册", authorized["answer"])
        self.assertIn("薪酬系统", authorized["answer"])

    def test_named_document_overview_is_fuzzy_acl_filtered_and_model_free(self):
        with tempfile.TemporaryDirectory() as root:
            project = Path(root)
            public_path = project / "document.md"
            public_path.write_text(
                "# MiniMax H3 提示词速查\n本手册介绍视频提示词。\n"
                "## 时间戳规则\n按时间顺序描述镜头。\n",
                encoding="utf-8",
            )
            private_path = project / "private.md"
            private_path.write_text(
                "# MiniMax 内部预算\n仅限财务人员。\n## 成本明细\n保密内容。\n",
                encoding="utf-8",
            )
            database = self.make_database(root)
            embeddings = EmbeddingClient(
                mode="hash", provider="builtin", model="hash-1024",
                base_url="", api_key="",
            )
            ingestion = self.make_ingestion(database, embeddings)
            ingestion.import_file(
                public_path, title="document", source_key="public-minimax",
                access_scope="public",
            )
            ingestion.import_file(
                private_path, title="private", source_key="private-minimax",
                access_scope="restricted",
            )
            service = ITQueryService(
                str(project / "missing.md"), database,
                ModelRouter("local", local_provider="ollama", local_model="qwen2.5:7b"),
                retriever=HybridRetriever(database, embeddings),
            )
            principal = Principal("viewer-1", frozenset({"viewer"}), frozenset())

            result = service.query("named-overview", "minmax讲了什么", principal)
            database.dispose()

        self.assertEqual(result["model"]["route"], "knowledge")
        self.assertIsNone(result["usage"])
        self.assertIn("MiniMax H3 提示词速查", result["answer"])
        self.assertIn("时间戳规则", result["answer"])
        self.assertNotIn("内部预算", result["answer"])
        self.assertEqual(result["citations"][0]["source"], "MiniMax H3 提示词速查")

    def test_password_protected_pdf_is_rejected_cleanly(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "protected.pdf"
            writer = PdfWriter()
            writer.add_blank_page(width=100, height=100)
            writer.encrypt("secret")
            with path.open("wb") as stream:
                writer.write(stream)
            with self.assertRaisesRegex(ValueError, "不支持加密 PDF"):
                parse_document(path, max_bytes=1024 * 1024)

    def test_provider_embeddings_are_batched_and_usage_is_attached_to_version(self):
        calls = 0
        request_ids = []

        def handler(request: httpx.Request):
            nonlocal calls
            calls += 1
            request_ids.append(request.headers.get("x-request-id"))
            payload = __import__("json").loads(request.content)
            vectors = [[0.0] * 1024 for _ in payload["input"]]
            for vector in vectors:
                vector[0] = 1.0
            return httpx.Response(200, headers={"x-request-id": f"embed-{calls}"}, json={
                "data": [
                    {"index": index, "embedding": vector}
                    for index, vector in enumerate(vectors)
                ],
                "usage": {"prompt_tokens": 12, "total_tokens": 12},
            })

        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "large.txt"
            path.write_text("第一段故障说明。\n" * 100, encoding="utf-8")
            database = self.make_database(root)
            embeddings = EmbeddingClient(
                mode="provider", provider="qwen", model="text-embedding-v3",
                base_url="https://embedding.test/v1", api_key="secret",
                batch_size=2, transport=httpx.MockTransport(handler),
            )
            result = self.make_ingestion(database, embeddings).import_file(
                path, source_key="manual/large",
            )
            version_id = result["version_id"]
            rows = database.document_usage_ledger(version_id)
            database.dispose()
        self.assertGreater(calls, 1)
        self.assertTrue(all(value and value.startswith("ingest-") for value in request_ids))
        self.assertEqual(len(rows), calls)
        self.assertTrue(all(row["operation"] == "embedding" for row in rows))
        self.assertTrue(all(row["query_id"] is None for row in rows))

    def test_embedding_failure_falls_back_to_lexical_search(self):
        with tempfile.TemporaryDirectory() as root:
            path = Path(root) / "dns.md"
            path.write_text("## DNS 缓存\n请清理 DNS 缓存并重新连接。", encoding="utf-8")
            database = self.make_database(root)
            hash_embeddings = EmbeddingClient(
                mode="hash", provider="builtin", model="hash-1024",
                base_url="", api_key="",
            )
            self.make_ingestion(database, hash_embeddings).import_file(path, source_key="dns")

            def failing(_request: httpx.Request):
                raise httpx.ConnectError("offline")

            failed_embeddings = EmbeddingClient(
                mode="provider", provider="qwen", model="text-embedding-v3",
                base_url="https://embedding.test/v1", api_key="secret",
                transport=httpx.MockTransport(failing),
            )
            query_id = database.record("fallback", "DNS 缓存", "pending", "pending")
            hits = HybridRetriever(database, failed_embeddings).retrieve("DNS 缓存", query_id)
            database.dispose()
        self.assertTrue(hits)
        self.assertIn("DNS 缓存", hits[0]["content"])


if __name__ == "__main__":
    unittest.main()
